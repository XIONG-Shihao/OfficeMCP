"""Word documents through python-docx.

Addresses: `body/p[N]` is the Nth paragraph directly in the body, `body/table[N]` the Nth table,
`body/table[N]/r[R]/c[C]` a cell and `.../c[C]/p[K]` a paragraph in it. Ordinals count from 1.
"""

import docx
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from ..catalog import EXPECTED_TEXT, POSITION, action, address, text
from ..errors import Refusal
from .common import excerpt, expect, pick, replace_in_runs, steps, clone

# Paragraph content other than plain runs: changing the paragraph's text could lose it.
RICH = {qn(t) for t in ("w:hyperlink", "w:fldSimple", "w:fldChar", "w:drawing", "w:object", "w:ins", "w:del",
                        "w:sdt", "w:footnoteReference", "w:endnoteReference", "w:commentRangeStart",
                        "w:commentReference", "w:bookmarkStart", "w:txbxContent", "w:sectPr")} | {
    "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath",
    "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent",
}


def body_blocks(document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def paragraphs(document):
    return [b for b in body_blocks(document) if isinstance(b, Paragraph)]


def tables(document):
    return [b for b in body_blocks(document) if isinstance(b, Table)]


def resolve(document, where: str):
    """The paragraph, table or cell an address names."""
    path = steps(where)
    if not path or path[0][0] != "body":
        raise Refusal("target_not_found", f"{where}: DOCX addresses start with body/", hint="For example body/p[3]")
    if len(path) == 1:
        return document
    kind, n = path[1]
    if kind == "p" and len(path) == 2:
        return pick(paragraphs(document), n, where, "paragraph")
    if kind != "table":
        raise Refusal("target_not_found", f"{where}: expected p[N] or table[N] after body/")
    table = pick(tables(document), n, where, "table")
    if len(path) == 2:
        return table
    if len(path) >= 4 and path[2][0] == "r" and path[3][0] == "c":
        rows = table.rows
        row = pick(list(rows), path[2][1], where, "row")
        cells = row.cells
        cell = pick(list(cells), path[3][1], where, "cell")
        if len(path) == 4:
            return cell
        if len(path) == 5 and path[4][0] == "p":
            return pick(cell.paragraphs, path[4][1], where, "paragraph in the cell")
    raise Refusal("target_not_found", f"{where} is not a DOCX address", hint="body/table[N]/r[R]/c[C][/p[K]]")


def expect_kind(target, kind, where: str):
    if not isinstance(target, kind):
        raise Refusal("target_not_found", f"{where} is not a {kind.__name__.strip('_').lower()}")
    return target


def plain(paragraph: Paragraph, where: str) -> None:
    """Refuse a paragraph whose text cannot be rewritten without losing non-text content."""
    for element in paragraph._p.iter():
        if element.tag in RICH:
            raise Refusal("unsupported_content", f"{where} holds {element.tag.split('}')[1]} content that a text change could lose",
                          hint="Use replace_text to change part of its text instead")


def set_paragraph(paragraph: Paragraph, value: str, where: str) -> None:
    plain(paragraph, where)
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(value)
        return
    runs[0].text = value
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


def cell_text(cell: _Cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs)


def heading_level(paragraph: Paragraph):
    name = paragraph.style.name if paragraph.style is not None else ""
    if name == "Title":
        return 0
    if name.startswith("Heading "):
        try:
            return int(name.split()[1])
        except (IndexError, ValueError):
            return None
    return None


def merged(table: Table) -> bool:
    return any(tc.find(qn("w:tcPr")) is not None and (tc.tcPr.find(qn("w:gridSpan")) is not None or tc.tcPr.find(qn("w:vMerge")) is not None)
               for tc in table._tbl.iter(qn("w:tc")))


def headers(table: Table, where: str) -> list[str]:
    if merged(table):
        raise Refusal("unsupported_content", f"{where} has merged cells", hint="Use set_cell_text on single cells")
    return [cell_text(c).strip() for c in table.rows[0].cells]


def column(names: list[str], name: str, where: str) -> int:
    if names.count(name) != 1:
        raise Refusal("ambiguous_match" if names.count(name) else "target_not_found",
                      f"{where}: the header '{name}' occurs {names.count(name)} times", headers=names)
    return names.index(name)


def set_cell(cell: _Cell, value: str, where: str) -> None:
    if cell._tc.find(qn("w:tbl")) is not None:
        raise Refusal("unsupported_content", f"{where} holds a nested table")
    paragraphs_ = cell.paragraphs
    for p in paragraphs_:
        plain(p, where)
    lines = value.split("\n")
    set_paragraph(paragraphs_[0], lines[0], where)
    for p in paragraphs_[1:]:
        p._p.getparent().remove(p._p)
    anchor = paragraphs_[0]
    for line in lines[1:]:
        new = clone(anchor._p)
        anchor._p.addnext(new)
        anchor = Paragraph(new, anchor._parent)
        set_paragraph(anchor, line, where)


def place(document, where: str, position: str, element) -> None:
    """Move a newly appended block element before or after the block an address names."""
    target = resolve(document, where)
    if isinstance(target, Paragraph):
        anchor = target._p
    elif isinstance(target, Table):
        anchor = target._tbl
    else:
        raise Refusal("target_not_found", f"{where}: insert next to a body paragraph or table")
    if anchor.getparent() is not document.element.body:
        raise Refusal("unsupported_content", f"{where}: insert next to blocks in the body, not inside a table")
    (anchor.addprevious if position == "before" else anchor.addnext)(element)


class Handler:
    format = "docx"
    extension = ".docx"

    def open(self, path):
        return docx.Document(str(path))

    def blank(self):
        return docx.Document()

    def save(self, document, path):
        document.save(str(path))

    def outline(self, document, limit=200):
        blocks, p_n, t_n = [], 0, 0
        for block in body_blocks(document):
            if isinstance(block, Paragraph):
                p_n += 1
                level = heading_level(block)
                item = {"address": f"body/p[{p_n}]", "kind": "heading" if level is not None else "paragraph",
                        "style": block.style.name if block.style is not None else None, "text": excerpt(block.text)}
                if level is not None:
                    item["level"] = level
            else:
                t_n += 1
                item = {"address": f"body/table[{t_n}]", "kind": "table", "rows": len(block.rows),
                        "columns": len(block.columns), "firstRow": [excerpt(cell_text(c), 60) for c in block.rows[0].cells] if block.rows else []}
            blocks.append(item)
        props = document.core_properties
        return {
            "paragraphs": p_n, "tables": t_n, "sections": len(document.sections),
            "properties": {k: getattr(props, k) for k in ("title", "subject", "author", "keywords") if getattr(props, k)},
            "blocks": blocks[:limit], "truncated": len(blocks) > limit,
        }

    def read(self, document, where, offset, limit):
        if not where:
            outline = self.outline(document, limit=100_000)
            blocks = outline["blocks"][offset: offset + limit]
            for b in blocks:
                if b["kind"] != "table":
                    b["text"] = resolve(document, b["address"]).text
            return {"blocks": blocks, "offset": offset, "next": offset + limit if offset + limit < len(outline["blocks"]) else None}
        target = resolve(document, where)
        if isinstance(target, Paragraph):
            return {"address": where, "kind": "paragraph", "style": target.style.name if target.style is not None else None,
                    "text": target.text,
                    "runs": [{"text": r.text, **{k: v for k, v in (("bold", r.bold), ("italic", r.italic), ("underline", r.underline)) if v}} for r in target.runs]}
        if isinstance(target, Table):
            return {"address": where, "kind": "table",
                    "rows": [[{"address": f"{where}/r[{i}]/c[{j}]", "text": cell_text(c)} for j, c in enumerate(row.cells, 1)]
                             for i, row in enumerate(target.rows, 1)][offset: offset + limit]}
        if isinstance(target, _Cell):
            return {"address": where, "kind": "cell", "text": cell_text(target),
                    "paragraphs": [{"address": f"{where}/p[{k}]", "text": p.text} for k, p in enumerate(target.paragraphs, 1)]}
        raise Refusal("target_not_found", f"{where} cannot be read")

    def find(self, document, needle, match_case, limit):
        found = []
        fold = (lambda s: s) if match_case else str.lower
        p_n, t_n = 0, 0
        for block in body_blocks(document):
            if isinstance(block, Paragraph):
                p_n += 1
                if fold(needle) in fold(block.text):
                    found.append({"address": f"body/p[{p_n}]", "text": excerpt(block.text), "count": fold(block.text).count(fold(needle))})
            else:
                t_n += 1
                for i, row in enumerate(block.rows, 1):
                    for j, cell in enumerate(row.cells, 1):
                        value = cell_text(cell)
                        if fold(needle) in fold(value):
                            found.append({"address": f"body/table[{t_n}]/r[{i}]/c[{j}]", "text": excerpt(value), "count": fold(value).count(fold(needle))})
            if len(found) >= limit:
                break
        return found[:limit]

    def text(self, document):
        out = []
        for block in body_blocks(document):
            if isinstance(block, Paragraph):
                out.append(block.text)
            else:
                out.extend("\t".join(cell_text(c) for c in row.cells) for row in block.rows)
        return "\n".join(out)


# ---- actions -------------------------------------------------------------------------------

PARAGRAPH = address("A paragraph: body/p[N], or body/table[N]/r[R]/c[C]/p[K].")
BLOCK = address("A body paragraph or table: body/p[N] or body/table[N].")
TABLE = address("A table: body/table[N].")
CELL = address("A table cell: body/table[N]/r[R]/c[C].")


@action("docx", "set_paragraph_text", "text", "Replace a paragraph's whole text, keeping its style and its first run's formatting.",
        {"target": PARAGRAPH, "expectedText": EXPECTED_TEXT, "text": text("The new text.")}, ["target", "expectedText", "text"])
def set_paragraph_text(document, args, edit):
    p = expect_kind(resolve(document, args["target"]), Paragraph, args["target"])
    expect(p.text, args["expectedText"], args["target"])
    set_paragraph(p, args["text"], args["target"])
    return {"changed": [args["target"]]}


@action("docx", "replace_text", "text",
        "Replace text wherever it occurs (in one target, or in the whole body and its tables), keeping each run's formatting.",
        {"find": text("The text to find.", 10_000), "replace": text("The replacement."),
         "target": address("Optional: limit to a paragraph, table or cell. Without it, the whole body and its tables."),
         "expectedCount": {"type": "integer", "minimum": 1, "description": "Refuse unless the text occurs exactly this many times."},
         "matchCase": {"type": "boolean", "default": True}},
        ["find", "replace"], tier="recipe")
def replace_text(document, args, edit):
    scope = resolve(document, args["target"]) if args.get("target") else document
    if isinstance(scope, Paragraph):
        candidates = [(args["target"], scope)]
    elif isinstance(scope, _Cell):
        candidates = [(f"{args['target']}/p[{k}]", p) for k, p in enumerate(scope.paragraphs, 1)]
    elif isinstance(scope, Table):
        candidates = [(f"{args['target']}/r[{i}]/c[{j}]/p[{k}]", p) for i, row in enumerate(scope.rows, 1)
                      for j, cell in enumerate(row.cells, 1) for k, p in enumerate(cell.paragraphs, 1)]
    else:
        candidates, t_n = [], 0
        for n, p in enumerate(paragraphs(document), 1):
            candidates.append((f"body/p[{n}]", p))
        for t_n, table in enumerate(tables(document), 1):
            for i, row in enumerate(table.rows, 1):
                seen = set()
                for j, cell in enumerate(row.cells, 1):
                    if id(cell._tc) in seen:  # a merged cell repeats across the columns it spans
                        continue
                    seen.add(id(cell._tc))
                    candidates.extend((f"body/table[{t_n}]/r[{i}]/c[{j}]/p[{k}]", p) for k, p in enumerate(cell.paragraphs, 1))
    match_case = args.get("matchCase", True)
    needle = args["find"] if match_case else args["find"].lower()
    hits = [(where, p) for where, p in candidates if needle in (p.text if match_case else p.text.lower())]
    total = sum((p.text if match_case else p.text.lower()).count(needle) for _, p in hits)
    if not total:
        raise Refusal("no_match", f"'{args['find']}' does not occur", hint="office_find shows where text occurs")
    if args.get("expectedCount") and total != args["expectedCount"]:
        raise Refusal("ambiguous_match", f"'{args['find']}' occurs {total} times, not {args['expectedCount']}", occurrences=total)
    changed = []
    for where, p in hits:
        if replace_in_runs(p.runs, args["find"], args["replace"], match_case):
            changed.append(where)
    # Text inside hyperlinks or fields is not in p.runs; refuse rather than half-replace.
    replacement = args["replace"] if match_case else args["replace"].lower()
    remaining = 0 if needle in replacement else sum((p.text if match_case else p.text.lower()).count(needle) for _, p in hits)
    if remaining:
        raise Refusal("unsupported_content", f"'{args['find']}' also occurs inside links or fields, which this action does not change")
    return {"changed": changed, "replaced": total}


@action("docx", "insert_paragraph", "structure", "Insert a paragraph before or after a body paragraph or table.",
        {"target": BLOCK, "position": POSITION, "text": text("The paragraph's text."),
         "style": {"type": "string", "maxLength": 200, "description": "A paragraph style the document defines, such as 'Normal' or 'List Bullet'."}},
        ["target", "text"])
def insert_paragraph(document, args, edit):
    style = args.get("style")
    if style and style not in [s.name for s in document.styles]:
        raise Refusal("invalid_request", f"The document has no style '{style}'", hint="office_inspect lists the styles in use")
    new = document.add_paragraph(args["text"], style)
    place(document, args["target"], args.get("position", "after"), new._p)
    return {"changed": [args["target"]], "inserted": "paragraph"}


@action("docx", "insert_heading", "structure", "Insert a heading (Word's Heading 1-9 styles) before or after a body paragraph or table.",
        {"target": BLOCK, "position": POSITION, "text": text("The heading's text.", 2000),
         "level": {"type": "integer", "minimum": 1, "maximum": 9, "description": "Heading level, 1 to 9."}},
        ["target", "text", "level"])
def insert_heading(document, args, edit):
    new = document.add_heading(args["text"], args["level"])
    place(document, args["target"], args.get("position", "after"), new._p)
    return {"changed": [args["target"]], "inserted": f"heading {args['level']}"}


@action("docx", "delete_paragraph", "structure", "Delete a body paragraph that holds only text.",
        {"target": address("A body paragraph: body/p[N]."), "expectedText": EXPECTED_TEXT}, ["target", "expectedText"])
def delete_paragraph(document, args, edit):
    p = expect_kind(resolve(document, args["target"]), Paragraph, args["target"])
    expect(p.text, args["expectedText"], args["target"])
    plain(p, args["target"])
    if len(paragraphs(document)) == 1 and not tables(document):
        raise Refusal("unsupported_content", "A document keeps at least one paragraph")
    p._p.getparent().remove(p._p)
    return {"changed": [args["target"]], "deleted": "paragraph"}


@action("docx", "set_paragraph_style", "formatting", "Apply a paragraph style the document defines.",
        {"target": PARAGRAPH, "style": {"type": "string", "minLength": 1, "maxLength": 200}}, ["target", "style"])
def set_paragraph_style(document, args, edit):
    p = expect_kind(resolve(document, args["target"]), Paragraph, args["target"])
    names = [s.name for s in document.styles]
    if args["style"] not in names:
        raise Refusal("invalid_request", f"The document has no style '{args['style']}'")
    p.style = document.styles[args["style"]]
    return {"changed": [args["target"]]}


@action("docx", "set_cell_text", "tables", "Replace a table cell's text (lines become paragraphs), keeping its first run's formatting.",
        {"target": CELL, "expectedText": EXPECTED_TEXT, "text": text("The new text; newlines start new paragraphs.")},
        ["target", "expectedText", "text"])
def set_cell_text(document, args, edit):
    cell = expect_kind(resolve(document, args["target"]), _Cell, args["target"])
    expect(cell_text(cell), args["expectedText"], args["target"])
    set_cell(cell, args["text"], args["target"])
    return {"changed": [args["target"]]}


@action("docx", "update_table_by_key", "tables",
        "In a table whose first row holds headers, find the one row whose key column has the key, and set one of its cells.",
        {"target": TABLE, "keyColumn": text("The header of the column that identifies rows.", 500), "key": text("The key value.", 2000),
         "column": text("The header of the column to change.", 500), "expectedText": text("Optional: the cell's current text."),
         "text": text("The new text.")},
        ["target", "keyColumn", "key", "column", "text"], tier="recipe")
def update_table_by_key(document, args, edit):
    table = expect_kind(resolve(document, args["target"]), Table, args["target"])
    names = headers(table, args["target"])
    key_index, target_index = column(names, args["keyColumn"], args["target"]), column(names, args["column"], args["target"])
    rows = [i for i, row in enumerate(table.rows[1:], 2) if cell_text(row.cells[key_index]).strip() == args["key"]]
    if len(rows) != 1:
        raise Refusal("ambiguous_match" if rows else "target_not_found", f"The key '{args['key']}' is in {len(rows)} rows")
    where = f"{args['target']}/r[{rows[0]}]/c[{target_index + 1}]"
    cell = table.rows[rows[0] - 1].cells[target_index]
    expect(cell_text(cell), args.get("expectedText"), where)
    set_cell(cell, args["text"], where)
    return {"changed": [where]}


@action("docx", "insert_table", "tables", "Insert a table of text before or after a body paragraph or table.",
        {"target": BLOCK, "position": POSITION,
         "rows": {"type": "array", "minItems": 1, "maxItems": 500, "items": {"type": "array", "minItems": 1, "maxItems": 50, "items": {"type": "string", "maxLength": 10000}},
                  "description": "Rows of cell texts; every row has the same number of cells."},
         "style": {"type": "string", "maxLength": 200, "description": "A table style the document defines; 'Table Grid' when it has one."}},
        ["target", "rows"])
def insert_table(document, args, edit):
    rows = args["rows"]
    if len({len(r) for r in rows}) != 1:
        raise Refusal("invalid_request", "Every row needs the same number of cells")
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    style = args.get("style") or ("Table Grid" if "Table Grid" in [s.name for s in document.styles] else None)
    if style:
        if style not in [s.name for s in document.styles]:
            raise Refusal("invalid_request", f"The document has no style '{style}'")
        table.style = document.styles[style]
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    place(document, args["target"], args.get("position", "after"), table._tbl)
    return {"changed": [args["target"]], "inserted": f"table {len(rows)}x{len(rows[0])}"}


@action("docx", "append_table_row", "tables", "Add a row at the end of a table, formatted like its last row.",
        {"target": TABLE, "cells": {"type": "array", "minItems": 1, "maxItems": 50, "items": {"type": "string", "maxLength": 10000},
                                    "description": "One text per column."}}, ["target", "cells"], tier="recipe")
def append_table_row(document, args, edit):
    table = expect_kind(resolve(document, args["target"]), Table, args["target"])
    if merged(table):
        raise Refusal("unsupported_content", f"{args['target']} has merged cells")
    last = table.rows[-1]
    if len(args["cells"]) != len(last.cells):
        raise Refusal("invalid_request", f"The table has {len(last.cells)} columns, not {len(args['cells'])}")
    for cell in last.cells:
        for p in cell.paragraphs:
            plain(p, args["target"])
    new = clone(last._tr)
    last._tr.addnext(new)
    row = table.rows[-1]
    for n, (cell, value) in enumerate(zip(row.cells, args["cells"]), 1):
        set_cell(cell, value, f"{args['target']}/r[{len(table.rows)}]/c[{n}]")
    return {"changed": [f"{args['target']}/r[{len(table.rows)}]"]}


@action("docx", "delete_table_row", "tables", "Delete one row of a table that has no merged cells.",
        {"target": address("A table row: body/table[N]/r[R]."),
         "expectedCells": {"type": "array", "items": {"type": "string"}, "description": "Optional: the row's current cell texts."}},
        ["target"])
def delete_table_row(document, args, edit):
    path = steps(args["target"])
    if len(path) != 3 or path[2][0] != "r":
        raise Refusal("target_not_found", f"{args['target']} is not a table row address")
    table = expect_kind(resolve(document, "/".join(f"{k}[{v}]" if v else k for k, v in path[:2])), Table, args["target"])
    if merged(table):
        raise Refusal("unsupported_content", f"{args['target']} is in a table with merged cells")
    row = pick(list(table.rows), path[2][1], args["target"], "row")
    current = [cell_text(c) for c in row.cells]
    if args.get("expectedCells") is not None and current != args["expectedCells"]:
        raise Refusal("expected_mismatch", f"{args['target']} does not hold the expected cells", actual=current)
    for cell in row.cells:
        for p in cell.paragraphs:
            plain(p, args["target"])
    if len(table.rows) == 1:
        raise Refusal("unsupported_content", "A table keeps at least one row")
    row._tr.getparent().remove(row._tr)
    return {"changed": [args["target"]], "deleted": "row"}


@action("docx", "add_comment", "review", "Attach a comment to a paragraph.",
        {"target": PARAGRAPH, "text": text("The comment.", 10_000), "author": {"type": "string", "maxLength": 200, "default": "OfficeMCP"}},
        ["target", "text"])
def add_comment(document, args, edit):
    p = expect_kind(resolve(document, args["target"]), Paragraph, args["target"])
    runs = p.runs or [p.add_run("")]
    document.add_comment(runs, text=args["text"], author=args.get("author", "OfficeMCP"))
    return {"changed": [args["target"]]}


@action("docx", "set_properties", "document", "Set the document's title, subject, author or keywords.",
        {k: {"type": "string", "maxLength": 1000} for k in ("title", "subject", "author", "keywords")}, [])
def set_properties(document, args, edit):
    props = document.core_properties
    changed = []
    for key in ("title", "subject", "author", "keywords"):
        if key in args:
            setattr(props, key, args[key])
            changed.append(key)
    if not changed:
        raise Refusal("invalid_request", "Name at least one of title, subject, author, keywords")
    return {"changed": [f"properties/{k}" for k in changed]}
