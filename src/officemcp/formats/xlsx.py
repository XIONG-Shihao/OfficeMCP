"""Excel workbooks through openpyxl.

Addresses: `Sales` (a sheet), `Sales!B2` (a cell), `Sales!B2:D9` (a range), `Sales!B:D` (columns)
and `table[Orders]` (an Excel table). A sheet name with spaces may be quoted: `'Q1 Sales'!A1`.
openpyxl does not calculate formulas: results say when Excel or LibreOffice must recalculate.
"""

import datetime as dt
import re

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table as ExcelTable

from ..catalog import CELL_VALUE, action, address, text
from ..errors import Refusal
from .common import excerpt

CELL = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")


class Book:
    """A workbook as openpyxl edits it, with the cached values Excel last calculated."""

    def __init__(self, workbook, values=None):
        self.workbook = workbook
        self.values = values  # data_only copy: cached results of formulas, when read from a file


def split(where: str):
    where = where.strip()
    if where.startswith("table[") and where.endswith("]"):
        return "table", where[6:-1], None
    sheet, bang, ref = where.rpartition("!")
    if not bang:
        return "sheet", where.strip("'"), None
    return "range", sheet.strip("'").replace("''", "'"), ref


def sheet_of(book: Book, name: str, where: str):
    if name not in book.workbook.sheetnames:
        raise Refusal("target_not_found", f"{where}: there is no sheet '{name}'", sheets=book.workbook.sheetnames)
    return book.workbook[name]


def find_table(book: Book, name: str, where: str):
    for ws in book.workbook.worksheets:
        if name in ws.tables:
            return ws, ws.tables[name]
    raise Refusal("target_not_found", f"{where}: there is no table '{name}'")


def bounds(ref: str, where: str):
    try:
        min_col, min_row, max_col, max_row = range_boundaries(ref.replace("$", ""))
    except Exception:
        raise Refusal("target_not_found", f"{where} is not a cell or range") from None
    if None in (min_col, max_col):
        raise Refusal("target_not_found", f"{where}: give rows and columns, for example B2:D9")
    return min_col, min_row or 1, max_col, max_row or 1


def region(book: Book, where: str, max_cells: int):
    """(worksheet, min_col, min_row, max_col, max_row) for a sheet, range or table address."""
    kind, name, ref = split(where)
    if kind == "table":
        ws, table = find_table(book, name, where)
        box = bounds(table.ref, where)
    else:
        ws = sheet_of(book, name, where)
        box = bounds(ref, where) if ref else (1, 1, max(ws.max_column, 1), max(ws.max_row, 1))
    cells = (box[2] - box[0] + 1) * (box[3] - box[1] + 1)
    if cells > max_cells:
        raise Refusal("too_large", f"{where} has {cells} cells, more than {max_cells} at once", hint="Read a smaller range")
    return (ws, *box)


def jsonable(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, dt.timedelta):
        return value.total_seconds()
    if isinstance(value, (int, float, str, bool)) or value is None:
        return value
    return str(value)


def current(cell):
    value = cell.value
    return jsonable(value)


def same(actual, expected) -> bool:
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)) and not isinstance(actual, bool):
        return abs(actual - expected) < 1e-9 * max(1, abs(expected))
    return actual == expected


def headers(ws, min_col, min_row, max_col):
    return [str(ws.cell(min_row, c).value).strip() if ws.cell(min_row, c).value is not None else "" for c in range(min_col, max_col + 1)]


def has_formulas(book: Book) -> bool:
    return any(isinstance(c.value, str) and c.value.startswith("=") for ws in book.workbook.worksheets for row in ws.iter_rows() for c in row)


def write(cell, value, edit):
    if cell.coordinate in cell.parent.merged_cells and not any(cell.coordinate == r.start_cell.coordinate for r in cell.parent.merged_cells.ranges):
        raise Refusal("unsupported_content", f"{cell.parent.title}!{cell.coordinate} is inside a merged cell")
    if isinstance(value, str) and value.startswith("="):
        edit.recalculation_needed = True
    cell.value = value


def references(book: Book, sheet: str) -> list[str]:
    """Formulas and defined names that name a sheet (renaming or deleting it would break them)."""
    pattern = re.compile(r"(^|[^A-Za-z0-9_'])('?)" + re.escape(sheet).replace("\\ ", " ") + r"\2!", re.IGNORECASE)
    found = []
    for ws in book.workbook.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and pattern.search(c.value):
                    found.append(f"{ws.title}!{c.coordinate}")
    for name, definition in book.workbook.defined_names.items():
        if pattern.search("=" + str(definition.attr_text)):
            found.append(f"name {name}")
    return found


class Handler:
    format = "xlsx"
    extension = ".xlsx"

    def open(self, path):
        try:
            workbook = openpyxl.load_workbook(path, keep_vba=False, keep_links=True, rich_text=True)
            values = openpyxl.load_workbook(path, data_only=True, read_only=True, keep_links=True)
        except Exception as error:
            raise Refusal("unreadable", f"Not a workbook openpyxl can open: {error}") from None
        return Book(workbook, values)

    def blank(self):
        return Book(openpyxl.Workbook())

    def save(self, book, path):
        book.workbook.save(str(path))

    def outline(self, book, limit=200):
        sheets = []
        for ws in book.workbook.worksheets[:limit]:
            sheets.append({
                "address": ws.title, "dimensions": ws.dimensions, "rows": ws.max_row, "columns": ws.max_column,
                "state": ws.sheet_state, "tables": [{"address": f"table[{t.displayName}]", "ref": t.ref} for t in ws.tables.values()],
                "mergedCells": len(ws.merged_cells.ranges),
            })
        return {"sheets": sheets, "definedNames": sorted(book.workbook.defined_names.keys()),
                "formulas": has_formulas(book), "truncated": len(book.workbook.worksheets) > limit}

    def read(self, book, where, offset, limit, max_cells=5000):
        if not where:
            return self.outline(book)
        ws, min_col, min_row, max_col, max_row = region(book, where, max_cells)
        cached = book.values[ws.title] if book.values is not None else None
        rows = []
        for r in range(min_row, max_row + 1):
            row = []
            for c in range(min_col, max_col + 1):
                cell = ws.cell(r, c)
                item = {"a": cell.coordinate, "v": current(cell)}
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    item["f"] = cell.value
                    item["v"] = jsonable(cached.cell(r, c).value) if cached is not None else None
                if cell.number_format and cell.number_format != "General":
                    item["format"] = cell.number_format
                row.append(item)
            rows.append(row)
        return {"address": where, "sheet": ws.title, "range": f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}",
                "cells": rows[offset: offset + limit], "next": offset + limit if offset + limit < len(rows) else None,
                "note": "v of a formula cell is the value Excel last saved; it may be stale after edits"}

    def find(self, book, needle, match_case, limit):
        fold = (lambda s: s) if match_case else str.lower
        found = []
        for ws in book.workbook.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None and fold(needle) in fold(str(c.value)):
                        found.append({"address": f"{ws.title}!{c.coordinate}", "text": excerpt(str(c.value))})
                        if len(found) >= limit:
                            return found
        return found

    def text(self, book):
        lines = []
        for ws in book.workbook.worksheets:
            lines.append(f"# {ws.title}")
            for row in ws.iter_rows(values_only=True):
                lines.append("\t".join("" if v is None else str(jsonable(v)) for v in row))
        return "\n".join(lines)


# ---- actions -------------------------------------------------------------------------------

GRID = {"type": "array", "minItems": 1, "maxItems": 10_000, "items": {"type": "array", "minItems": 1, "maxItems": 500, "items": CELL_VALUE},
        "description": "Rows of values, starting at the target's top-left cell."}
SHEET_OR_TABLE = address("A sheet name or table[Name].")


@action("xlsx", "set_cells", "cells", "Write values or formulas into a block of cells, starting at a cell.",
        {"target": address("The top-left cell: Sheet!B2."), "values": GRID,
         "expected": {**GRID, "description": "Optional: the current values (formulas as '=...'), same shape as values."}},
        ["target", "values"])
def set_cells(book, args, edit):
    kind, name, ref = split(args["target"])
    if kind != "range" or not CELL.match(ref or ""):
        raise Refusal("target_not_found", f"{args['target']} is not a cell like Sheet!B2")
    ws = sheet_of(book, name, args["target"])
    col, row = column_index_from_string(CELL.match(ref).group(1)), int(CELL.match(ref).group(2))
    values, expected = args["values"], args.get("expected")
    if expected is not None and [len(r) for r in expected] != [len(r) for r in values]:
        raise Refusal("invalid_request", "expected must have the same shape as values")
    for i, line in enumerate(values):
        for j, value in enumerate(line):
            cell = ws.cell(row + i, col + j)
            if expected is not None and not same(current(cell), expected[i][j]):
                raise Refusal("expected_mismatch", f"{ws.title}!{cell.coordinate} holds {current(cell)!r}, not {expected[i][j]!r}")
            write(cell, value, edit)
    if has_formulas(book):
        edit.recalculation_needed = True
    end = f"{get_column_letter(col + max(len(r) for r in values) - 1)}{row + len(values) - 1}"
    return {"changed": [f"{ws.title}!{get_column_letter(col)}{row}:{end}"]}


@action("xlsx", "clear_range", "cells", "Clear the values of a range, keeping its formatting.",
        {"target": address("A range: Sheet!B2:D9."), "expectedNonEmpty": {"type": "integer", "minimum": 0, "description": "Optional: how many cells hold values now."}},
        ["target"])
def clear_range(book, args, edit):
    ws, min_col, min_row, max_col, max_row = region(book, args["target"], edit.config.max_cells)
    cells = [ws.cell(r, c) for r in range(min_row, max_row + 1) for c in range(min_col, max_col + 1)]
    filled = [c for c in cells if c.value is not None]
    if args.get("expectedNonEmpty") is not None and len(filled) != args["expectedNonEmpty"]:
        raise Refusal("expected_mismatch", f"{args['target']} has {len(filled)} cells with values, not {args['expectedNonEmpty']}")
    for c in filled:
        write(c, None, edit)
    if has_formulas(book):
        edit.recalculation_needed = True
    return {"changed": [args["target"]], "cleared": len(filled)}


@action("xlsx", "set_number_format", "formatting", "Apply a number format (for example '0.00', '#,##0', 'yyyy-mm-dd') to a range.",
        {"target": address("A cell or range."), "format": {"type": "string", "minLength": 1, "maxLength": 255}}, ["target", "format"])
def set_number_format(book, args, edit):
    ws, min_col, min_row, max_col, max_row = region(book, args["target"], edit.config.max_cells)
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(r, c).number_format = args["format"]
    return {"changed": [args["target"]]}


@action("xlsx", "set_column_width", "formatting", "Set the width of one or more columns, in characters.",
        {"target": address("Columns: Sheet!B or Sheet!B:D."), "width": {"type": "number", "minimum": 0, "maximum": 255}}, ["target", "width"])
def set_column_width(book, args, edit):
    kind, name, ref = split(args["target"])
    if kind != "range" or not re.fullmatch(r"[A-Za-z]{1,3}(:[A-Za-z]{1,3})?", ref or ""):
        raise Refusal("target_not_found", f"{args['target']} is not a column like Sheet!B or Sheet!B:D")
    ws = sheet_of(book, name, args["target"])
    first, _, last = ref.upper().partition(":")
    for index in range(column_index_from_string(first), column_index_from_string(last or first) + 1):
        ws.column_dimensions[get_column_letter(index)].width = args["width"]
    return {"changed": [args["target"]]}


@action("xlsx", "append_rows", "rows",
        "Add rows after the last used row of a sheet, or at the end of a table (which grows to include them).",
        {"target": SHEET_OR_TABLE, "rows": GRID}, ["target", "rows"], tier="recipe")
def append_rows(book, args, edit):
    kind, name, _ = split(args["target"])
    if kind == "table":
        ws, table = find_table(book, name, args["target"])
        min_col, min_row, max_col, max_row = bounds(table.ref, args["target"])
        if table.totalsRowCount:
            raise Refusal("unsupported_content", f"{args['target']} has a totals row")
        if any(len(r) != max_col - min_col + 1 for r in args["rows"]):
            raise Refusal("invalid_request", f"{args['target']} has {max_col - min_col + 1} columns")
        below = [ws.cell(r, c) for r in range(max_row + 1, max_row + 1 + len(args["rows"])) for c in range(min_col, max_col + 1)]
        if any(c.value is not None for c in below):
            raise Refusal("unsupported_content", f"The cells below {args['target']} are not empty")
        start = max_row + 1
        table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row + len(args['rows'])}"
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
    elif kind == "sheet":
        ws = sheet_of(book, name, args["target"])
        min_col, start = 1, (ws.max_row + 1 if ws.max_row > 1 or ws.cell(1, 1).value is not None else 1)
    else:
        raise Refusal("target_not_found", f"{args['target']}: give a sheet name or table[Name]")
    for i, line in enumerate(args["rows"]):
        for j, value in enumerate(line):
            write(ws.cell(start + i, min_col + j), value, edit)
    if has_formulas(book):
        edit.recalculation_needed = True
    return {"changed": [f"{ws.title}!{start}:{start + len(args['rows']) - 1}"], "rows": len(args["rows"])}


@action("xlsx", "update_row_by_key", "rows",
        "In a sheet or table whose first row holds headers, find the one row whose key column has the key, and set cells by header.",
        {"target": SHEET_OR_TABLE, "keyColumn": text("The header of the column that identifies rows.", 500),
         "key": CELL_VALUE, "values": {"type": "object", "additionalProperties": CELL_VALUE, "minProperties": 1, "description": "New values by header."},
         "expected": {"type": "object", "additionalProperties": CELL_VALUE, "description": "Optional: current values by header."}},
        ["target", "keyColumn", "key", "values"], tier="recipe")
def update_row_by_key(book, args, edit):
    ws, min_col, min_row, max_col, max_row = region(book, args["target"], 10**9)
    names = headers(ws, min_col, min_row, max_col)

    def index(header):
        if names.count(header) != 1:
            raise Refusal("ambiguous_match" if names.count(header) else "target_not_found",
                          f"The header '{header}' occurs {names.count(header)} times", headers=names)
        return min_col + names.index(header)

    key_col = index(args["keyColumn"])
    rows = [r for r in range(min_row + 1, max_row + 1) if same(current(ws.cell(r, key_col)), args["key"])]
    if len(rows) != 1:
        raise Refusal("ambiguous_match" if rows else "target_not_found", f"The key {args['key']!r} is in {len(rows)} rows")
    changed = []
    for header, value in args.get("expected", {}).items():
        cell = ws.cell(rows[0], index(header))
        if not same(current(cell), value):
            raise Refusal("expected_mismatch", f"{ws.title}!{cell.coordinate} holds {current(cell)!r}, not {value!r}")
    for header, value in args["values"].items():
        cell = ws.cell(rows[0], index(header))
        write(cell, value, edit)
        changed.append(f"{ws.title}!{cell.coordinate}")
    if has_formulas(book):
        edit.recalculation_needed = True
    return {"changed": changed}


@action("xlsx", "add_sheet", "sheets", "Add an empty sheet.",
        {"name": {"type": "string", "minLength": 1, "maxLength": 31}, "after": address("Optional: the sheet to put it after; last otherwise.")},
        ["name"])
def add_sheet(book, args, edit):
    if args["name"] in book.workbook.sheetnames or re.search(r"[\[\]:*?/\\]", args["name"]):
        raise Refusal("invalid_request", f"'{args['name']}' is taken or not a valid sheet name")
    index = None
    if args.get("after"):
        index = book.workbook.sheetnames.index(sheet_of(book, args["after"], args["after"]).title) + 1
    book.workbook.create_sheet(args["name"], index)
    return {"changed": [args["name"]]}


@action("xlsx", "rename_sheet", "sheets", "Rename a sheet that no formula or defined name refers to.",
        {"target": address("The sheet."), "name": {"type": "string", "minLength": 1, "maxLength": 31}}, ["target", "name"])
def rename_sheet(book, args, edit):
    ws = sheet_of(book, args["target"], args["target"])
    if args["name"] in book.workbook.sheetnames or re.search(r"[\[\]:*?/\\]", args["name"]):
        raise Refusal("invalid_request", f"'{args['name']}' is taken or not a valid sheet name")
    used = references(book, ws.title)
    if used:
        raise Refusal("unsupported_content", f"{len(used)} formulas or names refer to '{ws.title}'; renaming would break them",
                      referencedBy=used[:20])
    ws.title = args["name"]
    return {"changed": [args["name"]]}


@action("xlsx", "delete_sheet", "sheets", "Delete a sheet that no formula or defined name refers to.",
        {"target": address("The sheet."), "expectedRows": {"type": "integer", "minimum": 0, "description": "Optional: its current number of used rows."}},
        ["target"])
def delete_sheet(book, args, edit):
    ws = sheet_of(book, args["target"], args["target"])
    if len(book.workbook.worksheets) == 1:
        raise Refusal("unsupported_content", "A workbook keeps at least one sheet")
    if args.get("expectedRows") is not None and ws.max_row != args["expectedRows"]:
        raise Refusal("expected_mismatch", f"{ws.title} has {ws.max_row} rows, not {args['expectedRows']}")
    used = references(book, ws.title)
    if used:
        raise Refusal("unsupported_content", f"{len(used)} formulas or names refer to '{ws.title}'", referencedBy=used[:20])
    if ws._images or ws._charts or ws.tables:
        raise Refusal("unsupported_content", f"{ws.title} has images, charts or tables", hint="Delete it in Excel")
    edit.allow("part:application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml")
    for key in ("feature:Excel merged cells", "feature:Excel data validations", "feature:Excel conditional formats", "feature:Excel hyperlinks"):
        edit.allow(key, 10**6)  # what lived only on the deleted sheet
    book.workbook.remove(ws)
    return {"changed": [args["target"]], "deleted": "sheet"}


@action("xlsx", "add_table", "tables", "Turn a range whose first row holds headers into an Excel table.",
        {"target": address("A range: Sheet!A1:D20."), "name": {"type": "string", "pattern": "^[A-Za-z_][A-Za-z0-9_.]{0,250}$"}},
        ["target", "name"])
def add_table(book, args, edit):
    kind, name, ref = split(args["target"])
    if kind != "range":
        raise Refusal("target_not_found", f"{args['target']} is not a range")
    ws = sheet_of(book, name, args["target"])
    min_col, min_row, max_col, max_row = bounds(ref, args["target"])
    if any(args["name"] in s.tables for s in book.workbook.worksheets):
        raise Refusal("invalid_request", f"A table named {args['name']} exists")
    names = headers(ws, min_col, min_row, max_col)
    if "" in names or len(set(names)) != len(names):
        raise Refusal("invalid_request", "The first row needs distinct, non-empty headers", headers=names)
    ws.add_table(ExcelTable(displayName=args["name"], ref=ref.replace("$", "")))
    return {"changed": [f"table[{args['name']}]"]}
