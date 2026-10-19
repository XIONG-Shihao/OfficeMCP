"""Regenerate redistributable synthetic producer fixtures; not needed to run tests.
Requires openpyxl and python-pptx. No proprietary documents or app dependencies.
"""
from pathlib import Path
import hashlib, json, importlib.metadata
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pptx.util import Inches
root=Path(__file__).parent/'corpus'
root.mkdir(exist_ok=True)
w=Workbook();s=w.active;s.title='Revenue'
s.append(['Item','Amount']);s.append(['Services',125]);s.append(['Hardware',25]);s['A4']='Total';s['B4']='=SUM(B2:B3)';s['B4'].number_format='$#,##0.00'
t=Table(displayName='Sales',ref='A1:B3');t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium9',showRowStripes=True);s.add_table(t)
chart=BarChart();chart.add_data(Reference(s,min_col=2,min_row=1,max_row=3),titles_from_data=True);s.add_chart(chart,'D1')
h=w.create_sheet('Hidden');h.sheet_state='hidden';h['A1']='Hidden producer evidence 77';w.save(root/'producer.xlsx')
p=Presentation();slide=p.slides.add_slide(p.slide_layouts[1]);slide.shapes.title.text='Producer fixture';slide.placeholders[1].text='Revenue 150\nMargin 42';slide.notes_slide.notes_text_frame.text='Independent producer note 77'
shape=slide.shapes.add_table(2,2, Inches(1), Inches(4), Inches(5), Inches(1));shape.table.cell(0,0).text='Measure';shape.table.cell(0,1).text='Value';shape.table.cell(1,0).text='Revenue';shape.table.cell(1,1).text='150'
p.slides.add_slide(p.slide_layouts[6]);p.save(root/'producer.pptx')
meta={'producers':{name:importlib.metadata.version(name) for name in ['openpyxl','python-pptx']},'content':'Synthetic, created for OfficeMCP tests; no third-party user documents.','sha256':{f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in root.glob('producer.*')}}
(root/'provenance.json').write_text(json.dumps(meta,indent=2)+'\n')
print(json.dumps(meta))
