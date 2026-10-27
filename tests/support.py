"""Test fixtures: real files of every format, made with the same libraries an author would use."""

import io
import shutil
import tempfile
import unittest
from pathlib import Path

from officemcp.files import Config
from officemcp.service import Service


class Case(unittest.TestCase):
    """A temporary project folder OfficeMCP may use, and a Service over it."""

    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix="officemcp-test-")).resolve()
        self.inputs = self.root / "inputs"
        self.inputs.mkdir()
        self.config = Config.from_dict({"roots": [str(self.inputs)]})
        self.service = Service(self.config)

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def ok(self, result):
        self.assertEqual(result["status"], "ok", result)
        return result

    def refused(self, result, code):
        self.assertEqual(result["status"], "refused", result)
        self.assertEqual(result["code"], code, result)
        return result

    def files(self):
        """The project's files, apart from the fixtures' images."""
        return sorted(p.name for p in self.inputs.iterdir() if p.suffix != ".png")

    def edit(self, path, actions, **options):
        revision = self.ok(self.service.inspect(str(path)))["revision"]
        return self.service.edit(str(path), revision, actions, **options)


def png_bytes(color=(200, 30, 30)) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (40, 30), color).save(buffer, "PNG")
    return buffer.getvalue()


def make_docx(folder: Path, name="report.docx") -> Path:
    import docx
    from docx.shared import Inches

    document = docx.Document()
    document.add_heading("Quarterly report", 1)
    p = document.add_paragraph()
    p.add_run("Revenue grew ").bold = True
    p.add_run("by 5% in the north region.")
    document.add_paragraph("The south region was flat.")
    table = document.add_table(rows=3, cols=3)
    for i, row in enumerate([["Region", "Revenue", "Owner"], ["North", "120", "Ana"], ["South", "80", "Bo"]]):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    image = folder / "logo.png"
    image.write_bytes(png_bytes())
    document.add_picture(str(image), width=Inches(0.5))
    document.add_paragraph("Contact: 中文 sales@example.com")
    path = folder / name
    document.save(path)
    return path


def make_xlsx(folder: Path, name="budget.xlsx") -> Path:
    import openpyxl
    from openpyxl.worksheet.table import Table

    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Budget"
    rows = [["Item", "Qty", "Price", "Total"], ["Paper", 10, 2.5, "=B2*C2"], ["Pens", 4, 1.25, "=B3*C3"]]
    for row in rows:
        ws.append(row)
    ws.add_table(Table(displayName="Lines", ref="A1:D3"))
    summary = workbook.create_sheet("Summary")
    summary["A1"] = "Grand total"
    summary["B1"] = "=SUM(Budget!D2:D3)"
    path = folder / name
    workbook.save(path)
    return path


def make_xlsx_with_chart(folder: Path, name="chart.xlsx") -> Path:
    """A workbook with a chart and a sparkline, as Excel-like producers write them (xlsxwriter)."""
    import xlsxwriter

    path = folder / name
    workbook = xlsxwriter.Workbook(str(path))
    ws = workbook.add_worksheet("Data")
    for i, value in enumerate([3, 5, 2, 8]):
        ws.write(i, 0, value)
    chart = workbook.add_chart({"type": "column"})
    chart.add_series({"values": "=Data!$A$1:$A$4"})
    ws.insert_chart("C2", chart)
    ws.add_sparkline("B1", {"range": "Data!A1:A4"})
    workbook.close()
    return path


def make_pptx(folder: Path, name="deck.pptx") -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Plan"
    body = slide.placeholders[1]
    body.text_frame.text = "Hire two engineers"
    body.text_frame.add_paragraph().text = "Ship in March"
    slide.notes_slide.notes_text_frame.text = "Mention the budget."
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Numbers"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for i, row in enumerate([["Quarter", "Revenue"], ["Q1", "120"]]):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    image = folder / "badge.png"
    image.write_bytes(png_bytes((20, 120, 200)))
    second.shapes.add_picture(str(image), Inches(6), Inches(1), width=Inches(1))
    path = folder / name
    presentation.save(path)
    return path


def make_pdf(folder: Path, name="form.pdf", pages=3, outline=True, form=True) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = folder / name
    c = canvas.Canvas(str(path), pagesize=A4)
    for n in range(1, pages + 1):
        c.drawString(72, 760, f"Page {n} of the contract")
        if outline:
            c.bookmarkPage(f"p{n}")
            if n == 1:
                c.addOutlineEntry("Start", "p1", level=0)
        if form and n == 1:
            c.acroForm.textfield(name="Customer", x=72, y=700, width=200, height=20, value="")
        c.showPage()
    c.save()
    return path
