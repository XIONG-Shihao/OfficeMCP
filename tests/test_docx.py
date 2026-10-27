import unittest

import docx

from .support import Case, make_docx


class DocxTests(Case):
    def setUp(self):
        super().setUp()
        self.path = make_docx(self.inputs)
        self.original = self.path.read_bytes()

    def reopen(self, result):
        return docx.Document(self.ok(result)["path"])

    def test_inspect_outlines_blocks_with_addresses_and_nothing_lost_on_save(self):
        result = self.ok(self.service.inspect(str(self.path)))
        blocks = result["outline"]["blocks"]
        self.assertEqual(blocks[0], {"address": "body/p[1]", "kind": "heading", "style": "Heading 1", "text": "Quarterly report", "level": 1})
        self.assertEqual(blocks[3]["address"], "body/table[1]")
        self.assertEqual(result["cannotKeep"], [])
        self.assertTrue(result["revision"].startswith("sha256:"))

    def test_read_and_find_return_addresses(self):
        cell = self.ok(self.service.read(str(self.path), "body/table[1]/r[2]/c[1]"))
        self.assertEqual(cell["text"], "North")
        found = self.ok(self.service.find(str(self.path), "south"))["matches"]
        self.assertEqual([m["address"] for m in found], ["body/p[3]", "body/table[1]/r[3]/c[1]"])

    def test_set_paragraph_text_keeps_style_and_first_run_formatting(self):
        document = self.reopen(self.edit(self.path, [{"action": "set_paragraph_text", "target": "body/p[2]",
                                                      "expectedText": "Revenue grew by 5% in the north region.", "text": "Revenue grew by 7%."}]))
        p = document.paragraphs[1]
        self.assertEqual(p.text, "Revenue grew by 7%.")
        self.assertTrue(p.runs[0].bold)

    def test_replace_text_spanning_runs_keeps_each_runs_formatting(self):
        result = self.ok(self.edit(self.path, [{"action": "replace_text", "find": "grew by", "replace": "rose by", "expectedCount": 1}]))
        self.assertEqual(result["receipts"][0]["replaced"], 1)
        p = docx.Document(result["path"]).paragraphs[1]
        self.assertEqual(p.text, "Revenue rose by 5% in the north region.")
        self.assertEqual([r.text for r in p.runs], ["Revenue rose by", " 5% in the north region."])
        self.assertTrue(p.runs[0].bold)

    def test_expected_text_mismatch_refuses_and_writes_nothing(self):
        result = self.refused(self.edit(self.path, [{"action": "set_paragraph_text", "target": "body/p[2]", "expectedText": "Something else", "text": "x"}]),
                              "expected_mismatch")
        self.assertEqual(result["details"]["index"], 0)
        self.assertEqual(self.path.read_bytes(), self.original)

    def test_a_batch_applies_all_or_nothing(self):
        result = self.edit(self.path, [{"action": "insert_heading", "target": "body/p[1]", "text": "Summary", "level": 2},
                                       {"action": "set_paragraph_text", "target": "body/p[99]", "expectedText": "", "text": "x"}])
        self.refused(result, "target_not_found")
        self.assertEqual(self.path.read_bytes(), self.original)

    def test_structure_tables_comments_and_properties(self):
        result = self.ok(self.edit(self.path, [
            {"action": "insert_heading", "target": "body/p[1]", "text": "Summary", "level": 2},
            {"action": "insert_paragraph", "target": "body/p[2]", "text": "Prepared by finance.", "style": "Normal"},
            {"action": "update_table_by_key", "target": "body/table[1]", "keyColumn": "Region", "key": "South", "column": "Revenue", "expectedText": "80", "text": "85"},
            {"action": "append_table_row", "target": "body/table[1]", "cells": ["West", "40", "Cy"]},
            {"action": "add_comment", "target": "body/p[4]", "text": "Check the north figure."},
            {"action": "set_properties", "title": "Q3"},
        ]))
        document = docx.Document(result["path"])
        self.assertEqual([p.text for p in document.paragraphs[:4]], ["Quarterly report", "Summary", "Prepared by finance.", "Revenue grew by 5% in the north region."])
        self.assertEqual(document.paragraphs[1].style.name, "Heading 2")
        self.assertEqual([c.text for c in document.tables[0].rows[2].cells], ["South", "85", "Bo"])
        self.assertEqual([c.text for c in document.tables[0].rows[3].cells], ["West", "40", "Cy"])
        self.assertEqual(len(list(document.comments)), 1)
        self.assertEqual(document.core_properties.title, "Q3")
        self.assertEqual(result["fidelity"], {"checked": True, "lost": []})

    def test_a_paragraph_with_a_picture_is_not_rewritten(self):
        outline = self.ok(self.service.inspect(str(self.path)))["outline"]["blocks"]
        picture = next(b for b in outline if b["kind"] == "paragraph" and b["text"] == "")
        self.refused(self.edit(self.path, [{"action": "set_paragraph_text", "target": picture["address"], "expectedText": "", "text": "x"}]),
                     "unsupported_content")
        self.refused(self.edit(self.path, [{"action": "delete_paragraph", "target": picture["address"], "expectedText": ""}]),
                     "unsupported_content")

    def test_ambiguous_key_and_wrong_count_refuse(self):
        self.refused(self.edit(self.path, [{"action": "replace_text", "find": "region", "replace": "area", "expectedCount": 1}]), "ambiguous_match")
        self.refused(self.edit(self.path, [{"action": "replace_text", "find": "absent", "replace": "x"}]), "no_match")


if __name__ == "__main__":
    unittest.main()
