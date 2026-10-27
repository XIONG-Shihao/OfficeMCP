import unittest

from pptx import Presentation

from .support import Case, make_pptx, png_bytes


class PptxTests(Case):
    def setUp(self):
        super().setUp()
        self.path = make_pptx(self.inputs)
        outline = self.ok(self.service.inspect(str(self.path)))["outline"]["outline"]
        self.title = next(s["address"] for s in outline[0]["shapes"] if s["kind"].startswith("placeholder:title"))
        self.body = next(s["address"] for s in outline[0]["shapes"] if s["text"].startswith("Hire"))
        self.table = next(s["address"] for s in outline[1]["shapes"] if s["kind"] == "table")

    def test_inspect_read_find(self):
        result = self.ok(self.service.inspect(str(self.path)))
        self.assertEqual(result["outline"]["slides"], 2)
        self.assertEqual(result["cannotKeep"], [])
        self.assertEqual(self.ok(self.service.read(str(self.path), "slide[1]/notes"))["text"], "Mention the budget.")
        found = self.ok(self.service.find(str(self.path), "q1"))["matches"]
        self.assertEqual(found, [{"address": f"{self.table}/r[2]/c[1]", "text": "Q1"}])

    def test_text_actions_keep_formatting_and_lines(self):
        result = self.ok(self.edit(self.path, [
            {"action": "set_shape_text", "target": self.body, "expectedText": "Hire two engineers\nShip in March", "text": "Hire three engineers\nShip in April\nReview in May"},
            {"action": "replace_text", "find": "Plan", "replace": "Roadmap", "expectedCount": 1},
            {"action": "set_shape_text", "target": f"{self.table}/r[2]/c[2]", "expectedText": "120", "text": "125"},
            {"action": "set_notes", "target": "slide[1]", "text": "Mention hiring.", "expectedText": "Mention the budget."},
        ]))
        deck = Presentation(result["path"])
        first = deck.slides[0]
        self.assertEqual(first.shapes.title.text_frame.text, "Roadmap")
        self.assertEqual([p.text for p in first.placeholders[1].text_frame.paragraphs], ["Hire three engineers", "Ship in April", "Review in May"])
        self.assertEqual(first.notes_slide.notes_text_frame.text, "Mention hiring.")
        table = next(s for s in deck.slides[1].shapes if s.has_table).table
        self.assertEqual(table.cell(1, 1).text, "125")

    def test_slides_are_added_moved_and_deleted(self):
        image = self.inputs / "dot.png"
        image.write_bytes(png_bytes((0, 0, 0)))
        result = self.ok(self.edit(self.path, [
            {"action": "add_slide", "layout": "Title and Content", "title": "Risks", "bullets": ["Hiring", "Scope"], "position": 2},
            {"action": "move_slide", "target": "slide[3]", "position": 1},
            {"action": "add_picture", "target": "slide[2]", "image": str(image), "left": 1, "top": 1, "width": 0.5},
        ]))
        deck = Presentation(result["path"])
        self.assertEqual([s.shapes.title.text_frame.text for s in deck.slides], ["Numbers", "Plan", "Risks"])
        deleted = self.ok(self.edit(self.path, [{"action": "delete_slide", "target": "slide[1]", "expectedTitle": "Numbers"}]))
        self.assertEqual([s.shapes.title.text_frame.text for s in Presentation(deleted["path"]).slides], ["Plan", "Risks"])
        self.assertEqual(deleted["fidelity"]["lost"], [])

    def test_wrong_expected_text_and_unknown_layout_refuse(self):
        self.refused(self.edit(self.path, [{"action": "set_shape_text", "target": self.title, "expectedText": "Other", "text": "x"}]), "expected_mismatch")
        self.refused(self.edit(self.path, [{"action": "add_slide", "layout": "Nope"}]), "target_not_found")
        self.refused(self.edit(self.path, [{"action": "delete_slide", "target": "slide[1]", "expectedTitle": "Numbers"}]), "expected_mismatch")


if __name__ == "__main__":
    unittest.main()
