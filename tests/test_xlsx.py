import unittest
from pathlib import Path

import openpyxl

from .support import Case, make_xlsx, make_xlsx_with_chart


class XlsxTests(Case):
    def setUp(self):
        super().setUp()
        self.path = make_xlsx(self.inputs)

    def test_inspect_and_read_give_formulas_and_cached_values(self):
        outline = self.ok(self.service.inspect(str(self.path)))["outline"]
        self.assertEqual([s["address"] for s in outline["sheets"]], ["Budget", "Summary"])
        self.assertEqual(outline["sheets"][0]["tables"], [{"address": "table[Lines]", "ref": "A1:D3"}])
        cells = self.ok(self.service.read(str(self.path), "Budget!D2:D3"))["cells"]
        self.assertEqual(cells[0][0]["f"], "=B2*C2")
        self.assertEqual(self.ok(self.service.read(str(self.path), "table[Lines]"))["range"], "A1:D3")

    def test_set_cells_with_expected_values_and_formulas(self):
        result = self.ok(self.edit(self.path, [{"action": "set_cells", "target": "Budget!B2", "values": [[12, 3]], "expected": [[10, 2.5]]},
                                               {"action": "set_cells", "target": "Budget!E1", "values": [["Tax"], ["=D2*0.1"]]}]))
        self.assertTrue(result["recalculationNeeded"])
        ws = openpyxl.load_workbook(result["path"])["Budget"]
        self.assertEqual((ws["B2"].value, ws["C2"].value, ws["E2"].value), (12, 3, "=D2*0.1"))
        self.refused(self.edit(self.path, [{"action": "set_cells", "target": "Budget!B2", "values": [[1]], "expected": [[99]]}]), "expected_mismatch")

    def test_append_rows_grows_the_table_and_update_row_by_key(self):
        result = self.ok(self.edit(self.path, [
            {"action": "append_rows", "target": "table[Lines]", "rows": [["Ink", 2, 7, "=B4*C4"]]},
            {"action": "update_row_by_key", "target": "table[Lines]", "keyColumn": "Item", "key": "Pens", "values": {"Qty": 6}, "expected": {"Qty": 4}},
        ]))
        wb = openpyxl.load_workbook(result["path"])
        self.assertEqual(wb["Budget"].tables["Lines"].ref, "A1:D4")
        self.assertEqual(wb["Budget"]["B3"].value, 6)
        self.refused(self.edit(self.path, [{"action": "update_row_by_key", "target": "Budget", "keyColumn": "Item", "key": "Nope", "values": {"Qty": 1}}]),
                     "target_not_found")

    def test_sheets_referenced_by_formulas_are_not_renamed_or_deleted(self):
        self.refused(self.edit(self.path, [{"action": "rename_sheet", "target": "Budget", "name": "Costs"}]), "unsupported_content")
        result = self.ok(self.edit(self.path, [{"action": "add_sheet", "name": "Notes"}, {"action": "rename_sheet", "target": "Notes", "name": "Memo"},
                                               {"action": "delete_sheet", "target": "Memo"}]))
        self.assertEqual(openpyxl.load_workbook(result["path"]).sheetnames, ["Budget", "Summary"])

    def test_formatting_and_new_tables(self):
        result = self.ok(self.edit(self.path, [{"action": "set_number_format", "target": "Budget!C2:C3", "format": "0.00"},
                                               {"action": "set_column_width", "target": "Budget!A:B", "width": 20},
                                               {"action": "set_cells", "target": "Summary!A3", "values": [["Name", "Score"], ["a", 1]]},
                                               {"action": "add_table", "target": "Summary!A3:B4", "name": "Scores"}]))
        wb = openpyxl.load_workbook(result["path"])
        self.assertEqual(wb["Budget"]["C2"].number_format, "0.00")
        self.assertEqual(wb["Budget"].column_dimensions["B"].width, 20)
        self.assertIn("Scores", wb["Summary"].tables)

    def test_a_loss_is_refused_until_the_model_allows_it(self):
        chart = make_xlsx_with_chart(self.inputs)
        original = chart.read_bytes()
        inspected = self.ok(self.service.inspect(str(chart)))
        self.assertTrue(inspected["cannotKeep"], "the probe save shows the loss before any edit")
        action = [{"action": "set_cells", "target": "Data!A1", "values": [[4]]}]
        result = self.refused(self.edit(chart, action), "fidelity_loss")
        self.assertTrue(result["details"]["lost"])
        self.assertIn("allowLoss", result["hint"])
        self.assertEqual(chart.read_bytes(), original, "nothing was written")
        saved = self.ok(self.edit(chart, action, allowLoss=True))
        self.assertEqual(saved["fidelity"]["lost"], result["details"]["lost"], "the result lists what was lost")
        self.assertEqual(saved["path"], str(chart))

    def test_csv_and_text_conversion(self):
        csv = self.ok(self.service.convert(str(self.path), "csv", sheet="Budget"))
        lines = Path(csv["path"]).read_text(encoding="utf-8").splitlines()
        self.assertEqual(lines[0], "Item,Qty,Price,Total")
        self.refused(self.service.convert(str(self.path), "pdf"), "engine_unavailable") if not self.config.libreoffice else None

    def test_create_a_workbook(self):
        result = self.ok(self.service.create("xlsx", "new.xlsx", [{"action": "set_cells", "target": "Sheet!A1", "values": [["a", 1]]}]))
        self.assertEqual(openpyxl.load_workbook(result["path"])["Sheet"]["B1"].value, 1)


if __name__ == "__main__":
    unittest.main()
